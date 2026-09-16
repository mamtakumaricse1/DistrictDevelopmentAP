from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "DDWMD-Keycloak-and-Microservices-Guide.docx"


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_styled(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x0B, 0x3D, 0x5C)
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_label_body(doc, label, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(label + " ")
    set_run_font(r1, bold=True, color=(0x0B, 0x3D, 0x5C))
    r2 = p.add_run(body)
    set_run_font(r2)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, bold=True, size=10, color=(0xFF, 0xFF, 0xFF))
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "0B3D5C")
        shd.set(qn("w:val"), "clear")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("District Works Dashboard")
    set_run_font(r, size=22, bold=True, color=(0x0B, 0x3D, 0x5C))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Keycloak, Realms, Login, and Microservices")
    set_run_font(r, size=16, bold=True, color=(0x1F, 0x6F, 0x8B))

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tag.add_run("A guide in simple words and in technical language")
    set_run_font(r, size=12, italic=True)
    add_para(
        doc,
        "This document is written for two readers at once: a Class 10 student who wants the idea, "
        "and a software developer who wants the exact moving parts in this project.",
        italic=True,
    )

    add_heading_styled(doc, "1. What this project is", 1)
    add_label_body(
        doc,
        "Simple:",
        "It is a website for a district office to watch development works. "
        "The first district in the data is Changlang, Arunachal Pradesh. Tirap is a second district used for testing. "
        "The software is the same for every district. Districts are rows of data, not a new app.",
    )
    add_label_body(
        doc,
        "Technical:",
        "This is the District Development Works Monitoring Dashboard (DDWMD). Phase 3 covers administration "
        "(districts, departments, agencies, users, master data, settings). Projects and dashboards come later. "
        "The browser SPA is @ddwmd/web. It never talks to Identity or Organization directly. "
        "It only calls the Gateway at /api/v1.",
    )

    add_heading_styled(doc, "2. The school-gate story (remember this picture)", 1)
    add_para(
        doc,
        "Imagine a school campus with one main gate and two office rooms.",
    )
    add_bullets(
        doc,
        [
            "The main gate is the Gateway. Visitors must enter here. The gate does not keep student records.",
            "The ID card office is Identity. It knows who you are, which class (role) you belong to, and whether your ID is real.",
            "The school office is Organization. It knows the school (district), departments, and other master lists.",
            "Keycloak is the government ID printer. It does not store your school timetable. It only proves “this person is really da.changlang”.",
            "A realm is like a separate ID office for one district. Changlang IDs are not mixed with Tirap IDs.",
        ],
    )
    add_label_body(
        doc,
        "Technical:",
        "This is a three-process backend: Gateway :3000 (no database), Identity :3001 (district_identity), "
        "Organization :3002 (district_organization). Keycloak :8080 is the OpenID Connect identity provider. "
        "One Keycloak server hosts many realms (system, changlang, tirap).",
    )

    add_heading_styled(doc, "3. What Keycloak is doing", 1)
    add_label_body(
        doc,
        "Simple:",
        "Keycloak is the login machine. It asks for username and password. This project never stores those passwords "
        "in PostgreSQL. After you type the password, Keycloak gives the website a short-lived pass (a token). "
        "The website shows that pass to our APIs. Our APIs check the pass, then decide what you may see.",
    )
    add_label_body(
        doc,
        "Technical:",
        "Keycloak is the IdP (identity provider). The SPA uses Authorization Code flow with PKCE (oidc-client-ts, "
        "client_id ddwmd-web). After login, Keycloak issues a JWT access token. The token’s iss claim is the realm "
        "issuer, for example http://localhost:8080/realms/changlang. Identity verifies signature against that realm’s JWKS, "
        "checks iss against registered issuers, then maps (issuer, email/sub) to the users table. Passwords never enter "
        "district_identity or district_organization.",
    )
    add_para(doc, "Keycloak does:")
    add_bullets(
        doc,
        [
            "Show the login page for the chosen district.",
            "Check username and password (and later MFA/LDAP if configured).",
            "Give back tokens (access token, and related OIDC tokens).",
            "Host one isolated user store per realm.",
        ],
    )
    add_para(doc, "Keycloak does not:")
    add_bullets(
        doc,
        [
            "Store districts, departments, or works data.",
            "Decide by itself that a Changlang officer may open Tirap records. Our APIs do that.",
            "Replace the Gateway. The browser still talks to port 3000 for /api/v1.",
        ],
    )

    add_heading_styled(doc, "4. What a realm is", 1)
    add_label_body(
        doc,
        "Simple:",
        "A realm is a locked cupboard of users. Changlang staff live in the changlang cupboard. "
        "Tirap staff live in the tirap cupboard. Super admins live in the system cupboard. "
        "A Changlang login cannot pretend to be Tirap, because the token says “I was issued by changlang”.",
    )
    add_label_body(
        doc,
        "Technical:",
        "A Keycloak realm is an isolation boundary: users, clients, roles, and later MFA/LDAP do not mix across realms. "
        "The JWT iss uniquely identifies the realm. Identity’s registered_issuers table (and Organization’s "
        "districts.keycloak_issuer) form the allow-list. Adding a district means: create a Keycloak realm + a districts row "
        "+ users. No new codebase.",
    )
    add_table(
        doc,
        ["Realm", "Who uses it", "Local issuer", "Example username"],
        [
            ["changlang", "Changlang officers", "http://localhost:8080/realms/changlang", "da.changlang"],
            ["tirap", "Tirap officers (test isolation)", "http://localhost:8080/realms/tirap", "da.tirap"],
            ["system", "SUPER_ADMIN only", "http://localhost:8080/realms/system", "sys.admin"],
        ],
    )
    add_para(
        doc,
        "All app-user Keycloak passwords are listed in the next section.",
        italic=True,
    )

    add_heading_styled(doc, "5. Local usernames and passwords (this PC only)", 1)
    add_label_body(
        doc,
        "Simple:",
        "These are practice keys for your laptop. They are not secret government passwords. "
        "Never use them on the internet or in production.",
    )
    add_label_body(
        doc,
        "Technical:",
        "Values come from .env.example, Prisma seeds, and deploy/keycloak/import/*.json. "
        "The postgres Windows superuser password is the one you typed when installing PostgreSQL — this project does not know it.",
    )

    add_heading_styled(doc, "5.1 PostgreSQL (the database)", 2)
    add_table(
        doc,
        ["What", "Value"],
        [
            ["Host (use this in psql)", "127.0.0.1"],
            ["Port", "5432"],
            ["App username", "dashboard"],
            ["App password", "dashboard"],
            ["Identity database", "district_identity"],
            ["Organization database", "district_organization"],
            ["Connection URL (Identity)", "postgresql://dashboard:dashboard@localhost:5432/district_identity?schema=public"],
            ["Connection URL (Organization)", "postgresql://dashboard:dashboard@localhost:5432/district_organization?schema=public"],
            ["Windows installer superuser", "postgres  (password = whatever you set during install)"],
        ],
    )
    add_para(
        doc,
        "In PowerShell, set $env:PGPASSWORD='dashboard' then connect with -h 127.0.0.1. "
        "Do not use localhost if Windows tries IPv6 and login fails.",
        italic=True,
    )

    add_heading_styled(doc, "5.2 Keycloak admin console (the ID printer’s back office)", 2)
    add_table(
        doc,
        ["What", "Value"],
        [
            ["URL", "http://localhost:8080"],
            ["Admin username", "admin"],
            ["Admin password", "admin"],
            ["Install folder (this PC)", "C:\\keycloak-26.7.3\\keycloak-26.7.3\\bin\\kc.bat"],
        ],
    )
    add_para(
        doc,
        "Use admin / admin only to look at realms. Daily dashboard login is on http://localhost:5173, not here.",
        italic=True,
    )

    add_heading_styled(doc, "5.3 Dashboard website login (Changlang, Tirap, System)", 2)
    add_para(doc, "Website: http://localhost:5173   — click the district card, then type the username below.")
    add_para(
        doc,
        "Password for every seeded Keycloak app user: ChangeMe!2026",
        bold=True,
    )
    add_table(
        doc,
        ["Click this on login page", "Username", "Email (also works)", "Password", "Role"],
        [
            ["Changlang", "da.changlang", "da.changlang@ddwmd.local", "ChangeMe!2026", "DISTRICT_ADMIN"],
            ["Changlang", "pwd.changlang", "pwd.changlang@ddwmd.local", "ChangeMe!2026", "DEPARTMENT_USER (PWD)"],
            ["Changlang", "viewer.changlang", "viewer.changlang@ddwmd.local", "ChangeMe!2026", "VIEWER"],
            ["Tirap", "da.tirap", "da.tirap@ddwmd.local", "ChangeMe!2026", "DISTRICT_ADMIN"],
            ["Tirap", "pwd.tirap", "pwd.tirap@ddwmd.local", "ChangeMe!2026", "DEPARTMENT_USER (PWD)"],
            ["System administration", "sys.admin", "sys.admin@ddwmd.local", "ChangeMe!2026", "SUPER_ADMIN"],
        ],
    )
    add_para(
        doc,
        "Changlang is the district you click, not a username. Do not type “changlang” in the username box.",
        italic=True,
    )

    add_heading_styled(doc, "5.4 Internal service key and app URLs", 2)
    add_table(
        doc,
        ["What", "Value"],
        [
            ["INTERNAL_API_KEY (services calling each other)", "dev-internal-key"],
            ["Website", "http://localhost:5173"],
            ["Gateway API", "http://localhost:3000/api/v1"],
            ["Identity API / Swagger", "http://localhost:3001/api/v1/docs"],
            ["Organization API / Swagger", "http://localhost:3002/api/v1/docs"],
            ["OIDC web client id", "ddwmd-web"],
            ["API audience in the JWT", "ddwmd-api"],
        ],
    )

    add_heading_styled(doc, "6. How to log in, and what happens at each click", 1)
    add_para(doc, "Use this URL only: http://localhost:5173  (not 127.0.0.1 — the redirect must match).")
    add_para(doc, "Step by step:")
    add_bullets(
        doc,
        [
            "You open the website. The login page asks Identity (through the Gateway) for GET /api/v1/auth/login-options. That list is public. You see Changlang, Tirap, and System administration.",
            "You click Changlang. You are not typing a “code”. Changlang is a district card. The website remembers the issuer in sessionStorage and sends you to Keycloak.",
            "Keycloak shows its own login form for realm changlang. Type username da.changlang and password ChangeMe!2026.",
            "Keycloak checks the password and redirects back to http://localhost:5173/auth/callback?code=...  That code is a one-time ticket, not your password.",
            "The website trades the code for tokens (PKCE). Then it calls GET /api/v1/auth/me with Authorization: Bearer <access_token>.",
            "Identity verifies the JWT, finds the user row, loads roles and permissions, and returns the profile. You land on the dashboard / administration screens.",
        ],
    )
    add_label_body(
        doc,
        "Technical login sequence:",
        "GET /auth/login-options (Identity reads Organization internal district-realms over HTTP) → "
        "signinRedirect to Keycloak authorize endpoint → user authenticates → redirect to /auth/callback → "
        "signinRedirectCallback (authorization code + PKCE verifier) → JWT in memory (oidc-client-ts + sessionStorage user store) → "
        "GET /auth/me → JwtAuthGuard + TokenVerifierService + IdentityUserDirectory.map.",
    )
    add_para(doc, "Test accounts (same password ChangeMe!2026):")
    add_table(
        doc,
        ["Click this", "Username", "Application role"],
        [
            ["Changlang", "da.changlang", "DISTRICT_ADMIN (Changlang only)"],
            ["Changlang", "pwd.changlang", "DEPARTMENT_USER (PWD, Changlang)"],
            ["Changlang", "viewer.changlang", "VIEWER"],
            ["Tirap", "da.tirap", "DISTRICT_ADMIN (Tirap only)"],
            ["System administration", "sys.admin", "SUPER_ADMIN (all districts)"],
        ],
    )

    add_heading_styled(doc, "7. What each piece is doing (microservices)", 1)

    add_heading_styled(doc, "7.1 Browser (the website)", 2)
    add_label_body(
        doc,
        "Simple:",
        "The pages you see: login, administration tabs. It hides buttons you should not use, but hiding is only for comfort. "
        "The real lock is on the server.",
    )
    add_label_body(
        doc,
        "Technical:",
        "@ddwmd/web, Vite port 5173. Proxies /api to Gateway :3000. Never authorize only in React. "
        "OIDC client ddwmd-web, redirect_uri {origin}/auth/callback.",
    )

    add_heading_styled(doc, "7.2 Gateway", 2)
    add_label_body(
        doc,
        "Simple:",
        "The receptionist. It does not keep files. If you ask for users, it walks to Identity. If you ask for districts, it walks to Organization. "
        "It also answers “are you alive?” by pinging both offices.",
    )
    add_label_body(
        doc,
        "Technical:",
        "@ddwmd/gateway on port 3000. No Prisma, no database. Routes /auth, /users, /roles, /permissions to Identity. "
        "Everything else (districts, departments, agencies, master-data, settings) to Organization. Forwards the Authorization header. "
        "GET /api/v1/health/ready aggregates Identity and Organization liveness. The gateway is not a policy enforcement point.",
    )

    add_heading_styled(doc, "7.3 Identity", 2)
    add_label_body(
        doc,
        "Simple:",
        "The ID office. Tables: users, roles, permissions, which user has which role, which issuers (realms) we trust.",
    )
    add_label_body(
        doc,
        "Technical:",
        "@ddwmd/identity on port 3001, database district_identity. Owns registered_issuers, users, roles, permissions, "
        "user_roles, user_departments, user_agencies. GET /auth/login-options, GET /auth/me, user CRUD, role listing. "
        "PUT /api/v1/internal/issuers is called by Organization when a district issuer changes (header x-internal-key).",
    )

    add_heading_styled(doc, "7.4 Organization", 2)
    add_label_body(
        doc,
        "Simple:",
        "The school office. Tables: districts, departments, agencies, master data, settings. It does not keep passwords or the full user list. "
        "When it needs “who is this token?”, it phones Identity’s /auth/me.",
    )
    add_label_body(
        doc,
        "Technical:",
        "@ddwmd/organization on port 3002, database district_organization. RemoteUserDirectory calls Identity GET /auth/me with the same Bearer token. "
        "GET /api/v1/internal/district-realms feeds Identity’s login picker. District UUIDs in Identity’s user_roles are copied IDs, not foreign keys across databases.",
    )

    add_heading_styled(doc, "7.5 Shared library (not a service)", 2)
    add_label_body(
        doc,
        "Simple:",
        "A toolbox both offices share: how to read a token, how to say “you may not enter Tirap”, how to write error JSON.",
    )
    add_label_body(
        doc,
        "Technical:",
        "@ddwmd/common: JWT verification, JwtAuthGuard, PermissionsGuard, AuthzService, error envelope, internal HTTP helpers. No Prisma models.",
    )

    add_heading_styled(doc, "7.6 PostgreSQL (two databases)", 2)
    add_label_body(
        doc,
        "Simple:",
        "One Postgres program on your PC, two notebooks. Identity’s notebook and Organization’s notebook. They do not share tables. "
        "That is why there are two databases.",
    )
    add_label_body(
        doc,
        "Technical:",
        "Same server, databases district_identity and district_organization, user dashboard. Schema files: "
        "apps/identity/prisma/schema.prisma and apps/organization/prisma/schema.prisma. pgAdmin or psql -h 127.0.0.1.",
    )

    add_heading_styled(doc, "8. How the three services talk", 1)
    add_para(doc, "They never JOIN each other’s tables. They call HTTP.")
    add_table(
        doc,
        ["From", "To", "Call", "Why"],
        [
            ["Browser", "Gateway :3000", "All /api/v1/…", "Single public door"],
            ["Gateway", "Identity :3001", "auth, users, roles, permissions", "People and login"],
            ["Gateway", "Organization :3002", "districts, departments, …", "Admin master data"],
            ["Identity", "Organization", "GET /internal/district-realms", "Build the login district list"],
            ["Organization", "Identity", "PUT /internal/issuers", "Keep JWT iss allow-list in sync"],
            ["Organization", "Identity", "GET /auth/me (Bearer)", "Map a token to AuthContext"],
        ],
    )
    add_label_body(
        doc,
        "Technical:",
        "Internal routes use x-internal-key (INTERNAL_API_KEY). Public routes use Bearer JWT. Each of Identity and Organization "
        "verifies the JWT itself. Authorization is scoped from AuthContext: a Changlang DISTRICT_ADMIN cannot read Tirap even with a guessed UUID.",
    )

    add_heading_styled(doc, "9. What runs on your computer (no Docker)", 1)
    add_table(
        doc,
        ["Window / program", "Address", "Job"],
        [
            ["PostgreSQL 18", "localhost:5432", "Two databases"],
            ["npm run dev:services", ":3001 :3002 :3000", "Identity, Organization, Gateway"],
            ["npm run dev:web", "http://localhost:5173", "Website"],
            ["kc.bat start-dev", "http://localhost:8080", "Keycloak + realms"],
        ],
    )
    add_para(
        doc,
        "Keycloak on this PC was installed at C:\\keycloak-26.7.3\\keycloak-26.7.3\\bin\\kc.bat. "
        "Realm JSON files live in D:\\Dashboard\\deploy\\keycloak\\import\\.",
        italic=True,
    )

    add_heading_styled(doc, "10. Authorization (after login)", 1)
    add_label_body(
        doc,
        "Simple:",
        "Login proves who you are. Permission proves what you may do. A viewer can look. A district admin can manage users in their district only. "
        "A system admin can create districts.",
    )
    add_label_body(
        doc,
        "Technical:",
        "Application RBAC is in PostgreSQL (roles, role_permissions, user_roles), not only Keycloak realm roles. "
        "AuthzService.assertPermission / assertDistrictAccess run inside Identity and Organization. "
        "SUPER_ADMIN comes from the system realm. Frontend menu hiding is UX only.",
    )

    add_heading_styled(doc, "11. Tiny glossary", 1)
    add_table(
        doc,
        ["Word", "Simple meaning", "Technical meaning"],
        [
            ["Token / JWT", "A signed hall pass", "JSON Web Token, signed by the realm, verified via JWKS"],
            ["Issuer (iss)", "Which ID office printed the pass", "Token claim; must match registered_issuers"],
            ["PKCE", "A secret so nobody steals the one-time code", "Proof Key for Code Exchange (S256)"],
            ["Realm", "A separate user cupboard", "Keycloak isolation unit"],
            ["Gateway", "Reception desk", "Reverse proxy, no DB"],
            ["Identity", "ID office", "Users, roles, issuers"],
            ["Organization", "School office", "Districts and master data"],
            ["AuthContext", "Who you are after mapping", "userId, roles, districtIds, permissions on the request"],
        ],
    )

    add_heading_styled(doc, "12. What is not built yet", 1)
    add_para(
        doc,
        "Project CRUD, progress photos, public dashboards, meetings, notifications, and reports are later phases. "
        "Do not look for those screens yet.",
    )

    add_heading_styled(doc, "13. If login fails (short checklist)", 1)
    add_bullets(
        doc,
        [
            "Open http://localhost:5173 not http://127.0.0.1:5173.",
            "Do not refresh a page whose address contains ?code= — that ticket is one-time.",
            "Username is da.changlang, not changlang.",
            "Keep three programs up: APIs, web, and Keycloak.",
            "Keycloak admin console: http://localhost:8080 (admin / admin) only to inspect realms, not for daily app login.",
        ],
    )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(18)
    r = footer.add_run(
        "End of guide. This matches the running local setup: Gateway, Identity, Organization, two PostgreSQL databases, "
        "and one Keycloak with realms system, changlang, and tirap."
    )
    set_run_font(r, size=10, italic=True, color=(0x44, 0x44, 0x44))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
